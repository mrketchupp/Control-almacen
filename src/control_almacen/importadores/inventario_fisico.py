"""Lectura del Excel de inventario físico (INVENTARIO DE REFACCIONAMIENTO...xlsx).

Solo lee; nunca modifica el archivo. Ver docs/06-formatos-excel.md, sección A.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from control_almacen.dominio import normalizar as n

HOJA_CATALOGO = "ARTICULOS_MX"
COLUMNAS = (
    "ITEM",
    "CODIGO AX",
    "DESCRIPCIÓN",
    "DIMENSION",
    "NP",
    "CANTIDAD",
    "UM",
    "CONSUMO",
    "INGRESO",
    "TOTAL",
)
_NOMBRE_HOJA = re.compile(r"CONTENEDOR\s*#?\s*(\d+)\s+(INVENTARIABLE|CONSUMIBLE)", re.IGNORECASE)


class ErrorFormato(ValueError):
    """El archivo no tiene la estructura esperada."""


@dataclass
class RenglonInventario:
    fila: int
    item: str | None
    codigo: int
    descripcion: str | None
    dimension: str | None
    np: str | None
    cantidad: Decimal
    um: str
    um_hoja: str | None
    consumo: Decimal | None
    ingreso: Decimal | None
    total: Decimal | None
    nota: str | None = None


@dataclass
class HojaInventario:
    nombre: str  # exacto, con espacios finales
    tabla: str | None
    contenedor: int
    clase: str  # INV / CONS
    orden: int
    fila_totales: int | None
    renglones: list[RenglonInventario] = field(default_factory=list)
    filas_vacias: list[int] = field(default_factory=list)

    @property
    def suma_cantidad(self) -> Decimal:
        return sum((r.cantidad for r in self.renglones), Decimal(0))

    @property
    def suma_total(self) -> Decimal:
        return sum((r.total or Decimal(0) for r in self.renglones), Decimal(0))


@dataclass
class LibroInventario:
    ruta: Path
    hojas: list[HojaInventario]
    catalogo: dict[int, str]


def _texto_nota(comentario) -> str | None:
    if comentario is None:
        return None
    texto = comentario.text or ""
    # Excel antepone "Autor:" a la nota; se conserva el contenido.
    return texto.strip() or None


def leer_inventario(ruta: Path) -> LibroInventario:
    ruta = Path(ruta)
    libro = load_workbook(ruta)  # fórmulas, tablas y notas
    valores = load_workbook(ruta, data_only=True)  # valores en caché
    hojas: list[HojaInventario] = []
    for orden, ws in enumerate(libro.worksheets, start=1):
        coincidencia = _NOMBRE_HOJA.search(ws.title)
        if not coincidencia or ws.sheet_state != "visible":
            continue
        contenedor = int(coincidencia.group(1))
        clase = "INV" if coincidencia.group(2).upper().startswith("INV") else "CONS"
        tabla = next(iter(ws.tables.values()), None)
        if tabla is not None:
            min_col, min_fila, max_col, max_fila = range_boundaries(tabla.ref)
            con_totales = bool(tabla.totalsRowCount)
        else:
            min_col, min_fila, max_col, max_fila = 1, 1, 10, ws.max_row
            con_totales = False
        encabezados = [n.mayusculas(c.value) for c in ws[min_fila]][min_col - 1 : max_col]
        indice = {nombre: encabezados.index(nombre) for nombre in COLUMNAS if nombre in encabezados}
        faltantes = [c for c in COLUMNAS if c not in indice]
        if faltantes:
            raise ErrorFormato(f"La hoja '{ws.title}' no tiene las columnas {faltantes}")
        ultima_datos = max_fila - 1 if con_totales else max_fila
        hoja = HojaInventario(
            nombre=ws.title,
            tabla=tabla.displayName if tabla is not None else None,
            contenedor=contenedor,
            clase=clase,
            orden=orden,
            fila_totales=max_fila if con_totales else None,
        )
        wv = valores[ws.title]
        for fila in range(min_fila + 1, ultima_datos + 1):
            celdas = [wv.cell(fila, min_col + i).value for i in range(len(encabezados))]
            valor = {nombre: celdas[i] for nombre, i in indice.items()}
            codigo = n.codigo_ax(valor["CODIGO AX"])
            if codigo is None:
                hoja.filas_vacias.append(fila)
                continue
            cantidad = n.decimal(valor["CANTIDAD"]) or Decimal(0)
            hoja.renglones.append(
                RenglonInventario(
                    fila=fila,
                    item=n.valor_a_texto(valor["ITEM"]),
                    codigo=codigo,
                    descripcion=n.valor_a_texto(valor["DESCRIPCIÓN"]),
                    dimension=n.valor_a_texto(valor["DIMENSION"]),
                    np=n.valor_a_texto(valor["NP"]),
                    cantidad=cantidad,
                    um=n.unidad(valor["UM"]),
                    um_hoja=n.valor_a_texto(valor["UM"]),
                    consumo=n.decimal(valor["CONSUMO"]),
                    ingreso=n.decimal(valor["INGRESO"]),
                    total=n.decimal(valor["TOTAL"]),
                    nota=_notas_de_fila(ws, fila, min_col, max_col),
                )
            )
        hojas.append(hoja)
    if not hojas:
        raise ErrorFormato("No se encontraron hojas 'CONTENEDOR #n INVENTARIABLE/CONSUMIBLE'.")
    catalogo = leer_catalogo_inventario(valores) if HOJA_CATALOGO in valores.sheetnames else {}
    return LibroInventario(ruta=ruta, hojas=hojas, catalogo=catalogo)


def _notas_de_fila(ws, fila: int, min_col: int, max_col: int) -> str | None:
    notas = [_texto_nota(ws.cell(fila, c).comment) for c in range(min_col, max_col + 1)]
    notas = [t for t in notas if t]
    return "\n".join(notas) or None


def leer_catalogo_inventario(libro) -> dict[int, str]:
    """Hoja oculta ARTICULOS_MX: código → producto."""
    catalogo: dict[int, str] = {}
    for codigo, descripcion, *_ in libro[HOJA_CATALOGO].iter_rows(values_only=True, max_col=2):
        codigo = n.codigo_ax(codigo)
        if codigo is not None and descripcion:
            catalogo[codigo] = n.valor_a_texto(descripcion)
    return catalogo
