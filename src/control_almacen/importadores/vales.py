"""Lectura del libro de vales (VALES DE SALIDA DLTA.xlsm): DIARIO, formularios y catálogo.

Solo lee; nunca modifica el archivo. Ver docs/06-formatos-excel.md, sección B.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from control_almacen.dominio import normalizar as n

HOJA_DIARIO = "DIARIO"
COLUMNAS_DIARIO = (
    "fecha", "folio", "pase_entrada", "pase_salida", "origen", "depto_origen", "destino",
    "depto_destino", "oc", "cantidad", "codigo", "descripcion", "clave", "um", "lote",
    "entrego", "recibio", "autorizo", "familia", "transferencia_consumo",
)  # fmt: skip
ENCABEZADOS_ESPERADOS = ("FECHA", "NO. FOLIO", "PASE DE ENTRADA", "PASE DE SALIDA")


class ErrorFormato(ValueError):
    """El archivo no tiene la estructura esperada."""


@dataclass
class RenglonDiario:
    """Un renglón del DIARIO tal como está, más sus valores interpretados."""

    fila: int
    crudo: tuple
    fecha: dt.date | None
    folio: int | None
    pase_entrada: str | None
    pase_salida: str | None
    origen: str | None
    depto_origen: str | None
    destino: str | None
    depto_destino: str | None
    oc: str | None
    cantidad: Decimal | None
    resto_cantidad: str | None
    codigo: int | None
    descripcion: str | None
    clave: str | None
    um: str | None
    lote: str | None
    entrego: str | None
    recibio: str | None
    autorizo: str | None
    familia: str | None
    transferencia_consumo: str | None
    errores: set[str] = field(default_factory=set)

    @property
    def perdido(self) -> bool:
        """Renglón sin folio ni código (p. ej., todo #REF!)."""
        return self.folio is None and self.codigo is None


def _texto_cero_nulo(valor: object) -> str | None:
    """Texto limpio; el 0 que dejaba la macro en campos vacíos se lee como vacío."""
    if isinstance(valor, int | float) and not isinstance(valor, bool) and valor == 0:
        return None
    texto = n.valor_a_texto(valor)
    return None if texto in (None, "0") else texto


def interpretar_renglon(fila: int, valores: tuple) -> RenglonDiario:
    crudo = tuple(valores[:20]) + (None,) * max(0, 20 - len(valores))
    v = dict(zip(COLUMNAS_DIARIO, crudo, strict=True))
    errores = {nombre for nombre, valor in v.items() if n.es_error_excel(valor)}
    cantidad, resto = n.separar_cantidad(v["cantidad"])
    return RenglonDiario(
        fila=fila,
        crudo=crudo,
        fecha=n.fecha(v["fecha"]),
        folio=n.codigo_ax(v["folio"]),
        pase_entrada=_texto_cero_nulo(v["pase_entrada"]),
        pase_salida=_texto_cero_nulo(v["pase_salida"]),
        origen=_texto_cero_nulo(v["origen"]),
        depto_origen=_texto_cero_nulo(v["depto_origen"]),
        destino=_texto_cero_nulo(v["destino"]),
        depto_destino=_texto_cero_nulo(v["depto_destino"]),
        oc=_texto_cero_nulo(v["oc"]),
        cantidad=cantidad,
        resto_cantidad=resto,
        codigo=n.codigo_ax(v["codigo"]) or None,
        descripcion=_texto_cero_nulo(v["descripcion"]),
        clave=_texto_cero_nulo(v["clave"]),
        um=_texto_cero_nulo(v["um"]),
        lote=_texto_cero_nulo(v["lote"]),
        entrego=n.nombre_persona(v["entrego"]),
        recibio=n.nombre_persona(v["recibio"]),
        autorizo=n.nombre_persona(v["autorizo"]),
        familia=_texto_cero_nulo(v["familia"]),
        transferencia_consumo=_texto_cero_nulo(v["transferencia_consumo"]),
        errores=errores,
    )


@dataclass
class PlantillaLeida:
    hoja: str
    origen: str | None
    depto_origen: str | None
    destino: str | None
    depto_destino: str | None
    entrega_nombre: str | None
    entrega_puesto: str | None
    recibe_nombre: str | None
    recibe_puesto: str | None
    autoriza_nombre: str | None
    observaciones: str | None


@dataclass
class LibroVales:
    ruta: Path
    renglones: list[RenglonDiario]
    plantillas: list[PlantillaLeida]
    catalogo: dict[int, str]


def leer_vales(ruta: Path) -> LibroVales:
    ruta = Path(ruta)
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if HOJA_DIARIO not in libro.sheetnames:
            raise ErrorFormato("El archivo no tiene la hoja DIARIO.")
        renglones = leer_diario(libro[HOJA_DIARIO])
        plantillas: list[PlantillaLeida] = []
        catalogo: dict[int, str] = {}
        for nombre in libro.sheetnames:
            if nombre == HOJA_DIARIO:
                continue
            ws = libro[nombre]
            filas = list(ws.iter_rows(min_row=1, max_row=62, max_col=11, values_only=True))
            if plantilla := leer_formulario(nombre, filas):
                plantillas.append(plantilla)
            if not catalogo:
                catalogo = leer_catalogo_formulario(ws)
    finally:
        libro.close()
    return LibroVales(ruta=ruta, renglones=renglones, plantillas=plantillas, catalogo=catalogo)


def leer_diario(ws) -> list[RenglonDiario]:
    filas = ws.iter_rows(min_row=1, max_col=20, values_only=True)
    encabezado = next(filas, None) or ()
    normalizados = tuple(n.mayusculas(x) or "" for x in encabezado[:4])
    if normalizados != ENCABEZADOS_ESPERADOS:
        raise ErrorFormato(f"Encabezados de DIARIO inesperados: {encabezado[:4]}")
    renglones = []
    for fila, valores in enumerate(filas, start=2):
        if any(x is not None for x in valores):
            renglones.append(interpretar_renglon(fila, valores))
    return renglones


def _celda(filas: list[tuple], referencia: str):
    columna = ord(referencia[0]) - ord("A")
    fila = int(referencia[1:]) - 1
    if fila >= len(filas) or columna >= len(filas[fila]):
        return None
    return filas[fila][columna]


def leer_formulario(hoja: str, filas: list[tuple]) -> PlantillaLeida | None:
    """Lee el encabezado y el bloque de firmas de una hoja-formulario.

    El bloque de firmas se localiza por etiquetas porque no está en la misma fila
    en todas las hojas (en MECANICO y OPERACION DIA está desplazado).
    """
    if n.mayusculas(_celda(filas, "C17")) is None or "ORIGEN" not in (
        n.mayusculas(_celda(filas, "C17")) or ""
    ):
        return None
    fila_nombre = None
    for i in range(38, min(len(filas), 62)):
        if (n.mayusculas(filas[i][2] if len(filas[i]) > 2 else None) or "").startswith("NOMBRE"):
            fila_nombre = i + 1
            break
    autoriza = None
    for i in range(40, min(len(filas), 62)):
        fila = filas[i]
        if len(fila) > 6 and (n.mayusculas(fila[6]) or "").startswith("AUTORIZA"):
            for j in range(i + 1, min(len(filas), i + 4)):
                if len(filas[j]) > 6 and (nombre := n.nombre_persona(filas[j][6])):
                    autoriza = nombre
                    break
    observaciones = []
    for i in range(41, min(len(filas), 50)):
        texto = n.valor_a_texto(filas[i][2] if len(filas[i]) > 2 else None)
        etiquetas = ("OBSERVACION", "ENTREGO", "FIRMA", "NOMBRE", "PUESTO")
        if texto and not texto.upper().startswith(etiquetas) and not texto.isdigit():
            observaciones.append(texto)
    return PlantillaLeida(
        hoja=hoja,
        origen=n.valor_a_texto(_celda(filas, "E17")),
        depto_origen=n.valor_a_texto(_celda(filas, "I17")),
        destino=n.valor_a_texto(_celda(filas, "E18")),
        depto_destino=n.valor_a_texto(_celda(filas, "I18")),
        entrega_nombre=n.nombre_persona(_celda(filas, f"D{fila_nombre}")) if fila_nombre else None,
        entrega_puesto=n.mayusculas(_celda(filas, f"D{fila_nombre + 1}")) if fila_nombre else None,
        recibe_nombre=n.nombre_persona(_celda(filas, f"I{fila_nombre}")) if fila_nombre else None,
        recibe_puesto=n.mayusculas(_celda(filas, f"I{fila_nombre + 1}")) if fila_nombre else None,
        autoriza_nombre=autoriza,
        observaciones="\n".join(observaciones) or None,
    )


def leer_catalogo_formulario(ws) -> dict[int, str]:
    """Catálogo copiado en las columnas AG:AH de las hojas-formulario."""
    catalogo: dict[int, str] = {}
    for codigo, descripcion in ws.iter_rows(min_row=6, min_col=33, max_col=34, values_only=True):
        codigo = n.codigo_ax(codigo)
        if codigo is not None and descripcion:
            catalogo[codigo] = n.valor_a_texto(descripcion)
    return catalogo
