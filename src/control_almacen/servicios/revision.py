"""Lista de revisión del historial (formato v2): se genera desde los archivos del usuario,
el usuario la contesta en Excel y la herramienta aplica sus respuestas en la primera carga.

Es un archivo NUEVO (no una plantilla del usuario), por eso se crea con openpyxl.
"""

from __future__ import annotations

import datetime as dt
import difflib
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from control_almacen.dominio import normalizar as n
from control_almacen.dominio.catalogo import es_error_de_dedo
from control_almacen.importadores.inventario_fisico import LibroInventario
from control_almacen.importadores.vales import LibroVales, RenglonDiario
from control_almacen.servicios.limpieza import VACIO, RespuestasRevision

MARCA_FORMATO = "REVISION-V2"
SI, NO = "Sí", "No"

HOJA_RENGLONES = "1 Renglones a corregir"
HOJA_DUPLICADOS = "2 Duplicados"
HOJA_FOLIOS = "3 Folios a verificar"
HOJA_CODIGOS = "4 Códigos fuera de catálogo"
HOJA_NOMBRES = "5 Nombres a unificar"
HOJA_VALORES = "6 Valores a normalizar"
HOJA_INVENTARIO = "7 Inventario físico"

# Columnas editables de la hoja 1: (encabezado, campo de RenglonDiario)
EDITABLES = [
    ("Fecha", "fecha"),
    ("Origen", "origen"),
    ("Depto. origen", "depto_origen"),
    ("Destino", "destino"),
    ("Depto. destino", "depto_destino"),
    ("O.C.", "oc"),
    ("Cantidad", "cantidad"),
    ("Código", "codigo"),
    ("Descripción", "descripcion"),
    ("Clave", "clave"),
    ("U.M.", "um"),
    ("Lote", "lote"),
    ("Entregó", "entrego"),
    ("Recibió", "recibio"),
    ("Autorizó", "autorizo"),
]
FIJAS_HOJA1 = ["Prioridad", "Fila en DIARIO", "Folio", "Situación"]
CAMPOS_VALORES = {
    "U.M.": "um",
    "Lote / C.U": "lote",
    "Origen": "origen",
    "Depto. origen": "depto_origen",
    "Destino": "destino",
    "Depto. destino": "depto_destino",
}
SINONIMOS = {
    "um": {
        "CUB": "CUBETA",
        "LITROS": "LTS",
        "LT": "LTS",
        "M": "MTS",
        "MT": "MTS",
        "PZAS": "PZA",
        "PZ": "PZA",
    },
    "lote": {"NEUVO": "NUEVO", "NUEV0": "NUEVO", "NUEVA": "NUEVO", "O": VACIO},
}

_FUENTE = "Arial"
_ENCABEZADO = PatternFill("solid", fgColor="1F3864")
_EDITABLE = PatternFill("solid", fgColor="FFF2CC")
_PERDIDO = PatternFill("solid", fgColor="F8CBAD")
_PRIORIDAD = {
    "Alta": PatternFill("solid", fgColor="F8CBAD"),
    "Media": PatternFill("solid", fgColor="FFE699"),
    "Baja": PatternFill("solid", fgColor="E2EFDA"),
}
_BORDE = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


class ErrorRevision(ValueError):
    """El archivo no es una lista de revisión v2."""


# ======================================================================
# Detección
# ======================================================================


@dataclass
class Hallazgo:
    renglon: RenglonDiario
    prioridad: str
    situacion: str
    eliminar: bool = False


def _por_folio(renglones: list[RenglonDiario]) -> dict[int, list[RenglonDiario]]:
    grupos: dict[int, list[RenglonDiario]] = defaultdict(list)
    for r in renglones:
        if r.folio is not None:
            grupos[r.folio].append(r)
    return grupos


def _fechas_folio(grupos: dict[int, list[RenglonDiario]]) -> dict[int, dt.date]:
    fechas = {}
    for folio, filas in grupos.items():
        validas = [r.fecha for r in filas if r.fecha]
        if validas:
            fechas[folio] = Counter(validas).most_common(1)[0][0]
    return fechas


def detectar_renglones(renglones: list[RenglonDiario], catalogo: dict[int, str]) -> list[Hallazgo]:
    grupos = _por_folio(renglones)
    fechas = _fechas_folio(grupos)
    folios = sorted(fechas)
    hallazgos = []
    for r in renglones:
        motivos, prioridad, eliminar = [], None, False
        if r.perdido:
            motivos.append(
                "Renglón completamente perdido (#REF!). Si el PDF no muestra un renglón faltante, elimínalo."
            )
            prioridad, eliminar = "Alta", True
        else:
            criticos = r.errores & {
                "fecha",
                "folio",
                "cantidad",
                "codigo",
                "origen",
                "destino",
                "depto_origen",
                "depto_destino",
            }
            if criticos:
                motivos.append("Datos perdidos (#REF!): " + ", ".join(sorted(criticos)))
                prioridad = "Alta"
            otros = r.errores - criticos - {"pase_entrada", "pase_salida", "lote"}
            if otros:
                motivos.append("Datos perdidos (#REF!): " + ", ".join(sorted(otros)))
                prioridad = prioridad or "Media"
            if r.errores == {"lote"} or (
                r.errores and r.errores <= {"lote", "pase_entrada", "pase_salida"}
            ):
                motivos.append(
                    "Solo se perdió el LOTE / C.U. Si el PDF no dice nada, déjalo vacío."
                )
                prioridad = prioridad or "Baja"
            if isinstance(r.crudo[0], str) and r.fecha:
                motivos.append("Fecha escrita como texto (se convertirá)")
                prioridad = prioridad or "Baja"
            if r.fecha and r.folio in fechas and folios:
                i = folios.index(r.folio)
                vecinas = [fechas[f] for f in folios[max(0, i - 1) : i + 2]]
                if r.fecha < min(vecinas) - dt.timedelta(days=20) or r.fecha > max(
                    vecinas
                ) + dt.timedelta(days=20):
                    motivos.append(
                        f"Fecha {r.fecha:%d/%m/%Y} fuera de rango respecto a los folios vecinos"
                    )
                    prioridad = prioridad or "Media"
            if r.cantidad is None or r.cantidad <= 0 or r.resto_cantidad:
                motivos.append(f"Cantidad inválida: {r.crudo[9]!r}")
                prioridad = prioridad or "Media"
            if r.codigo is None and "codigo" not in r.errores:
                motivos.append("Sin código")
                prioridad = prioridad or "Media"
            oficial = catalogo.get(r.codigo) if r.codigo else None
            if (
                oficial
                and r.descripcion
                and not es_error_de_dedo(r.descripcion, oficial)
                and n.clave_laxa(r.descripcion) != n.clave_laxa(oficial)
            ):
                sugerido = next(
                    (
                        c
                        for c, d in catalogo.items()
                        if n.clave_laxa(d) == n.clave_laxa(r.descripcion) and c != r.codigo
                    ),
                    None,
                )
                texto = f"La descripción '{r.descripcion}' no corresponde al código {r.codigo} ({oficial})"
                if sugerido:
                    texto += f". ¿Debía ser el código {sugerido}?"
                motivos.append(texto)
                prioridad = prioridad or "Media"
        if motivos:
            hallazgos.append(Hallazgo(r, prioridad or "Baja", " · ".join(motivos), eliminar))
    orden = {"Alta": 0, "Media": 1, "Baja": 2}
    return sorted(hallazgos, key=lambda h: (orden[h.prioridad], h.renglon.fila))


def sugerencias_encabezado(renglones: list[RenglonDiario]) -> dict[int, dict[str, object]]:
    """Para cada fila con encabezado perdido: el valor que tienen los otros renglones del folio."""
    grupos = _por_folio(renglones)
    sugerencias: dict[int, dict[str, object]] = {}
    for filas in grupos.values():
        for r in filas:
            for _, campo in EDITABLES:
                if campo in r.errores or getattr(r, campo) is None:
                    otro = next(
                        (
                            getattr(x, campo)
                            for x in filas
                            if x is not r
                            and campo not in x.errores
                            and getattr(x, campo) is not None
                        ),
                        None,
                    )
                    if otro is not None and campo in (
                        "fecha",
                        "origen",
                        "depto_origen",
                        "destino",
                        "depto_destino",
                        "entrego",
                        "recibio",
                        "autorizo",
                    ):
                        sugerencias.setdefault(r.fila, {})[campo] = otro
    return sugerencias


def detectar_duplicados(
    renglones: list[RenglonDiario],
) -> list[tuple[RenglonDiario, RenglonDiario, str]]:
    grupos = _por_folio(renglones)
    resultado = []
    for folio in sorted(grupos):
        filas = grupos[folio]
        vistos: dict[tuple, RenglonDiario] = {}
        pares = []
        for r in filas:
            clave = r.crudo[:18]
            if clave in vistos:
                pares.append((vistos[clave], r))
            else:
                vistos[clave] = r
        if pares:
            completo = len(filas) == 2 * len(pares)
            tipo = (
                "Vale completo guardado 2 veces" if completo else "Renglón repetido dentro del vale"
            )
            resultado.extend((a, b, tipo) for a, b in pares)
    return resultado


def detectar_folios(renglones: list[RenglonDiario]) -> list[tuple[str, str, str, str]]:
    """(prioridad, folio, situación, qué revisar) — solo informativo."""
    grupos = _por_folio(renglones)
    salida = []
    if grupos:
        todos = sorted(grupos)
        for faltante in (f for f in range(todos[0], todos[-1] + 1) if f not in grupos):
            salida.append(
                (
                    "Alta",
                    str(faltante),
                    "Folio que no aparece en DIARIO",
                    "¿Existe el PDF? Si existe, sus renglones se capturarán en la herramienta.",
                )
            )
    for folio, filas in sorted(grupos.items()):
        if len(filas) >= 16:
            salida.append(
                (
                    "Alta",
                    str(folio),
                    f"{len(filas)} renglones: pudo perder renglones al guardar (límite de 16–18) o se guardó dos veces",
                    "Contar los renglones del PDF.",
                )
            )
        deptos = {n.compactar(r.depto_destino) for r in filas if r.depto_destino}
        if len(deptos) > 1:
            salida.append(
                (
                    "Media",
                    str(folio),
                    "El folio tiene más de un departamento destino: " + " / ".join(sorted(deptos)),
                    "¿Son dos vales distintos con el mismo número?",
                )
            )
    return salida


def _agrupar_variantes(valores: Counter, umbral: float = 0.88) -> list[list[tuple[str, int]]]:
    nombres = sorted(valores, key=lambda k: -valores[k])
    usados: set[str] = set()
    grupos = []
    for a in nombres:
        if a in usados:
            continue
        grupo = [a]
        usados.add(a)
        for b in nombres:
            if b in usados:
                continue
            la, lb = n.clave_laxa(a), n.clave_laxa(b)
            parecidos = la == lb or difflib.SequenceMatcher(None, la, lb).ratio() >= umbral
            prefijo = len(lb) > 8 and (la.startswith(lb) or lb.startswith(la))
            if parecidos or prefijo:
                grupo.append(b)
                usados.add(b)
        if len(grupo) > 1:
            grupos.append([(x, valores[x]) for x in grupo])
    return grupos


def detectar_nombres(renglones: list[RenglonDiario]) -> list[tuple[int, str, int, str]]:
    """(grupo, variante, veces, nombre propuesto)."""
    conteo = Counter(x for r in renglones for x in (r.entrego, r.recibio, r.autorizo) if x)
    salida = []
    for i, grupo in enumerate(_agrupar_variantes(conteo, 0.86), start=1):
        propuesto = sorted(grupo, key=lambda x: (-len(x[0].split()), -x[1]))[0][0]
        salida.extend((i, nombre, veces, propuesto) for nombre, veces in grupo)
    return salida


def detectar_valores(renglones: list[RenglonDiario]) -> list[tuple[str, str, int, str]]:
    """(campo, valor encontrado, renglones, valor propuesto)."""
    salida = []
    for etiqueta, campo in CAMPOS_VALORES.items():
        conteo = Counter(
            (n.unidad(getattr(r, campo)) if campo == "um" else getattr(r, campo)) or VACIO
            for r in renglones
            if not r.perdido and campo not in r.errores
        )
        propuestas: dict[str, str] = {}
        for valor in conteo:
            if valor in SINONIMOS.get(campo, {}):
                propuestas[valor] = SINONIMOS[campo][valor]
        for grupo in _agrupar_variantes(Counter({k: v for k, v in conteo.items() if k != VACIO})):
            canonico = grupo[0][0]
            for valor, _ in grupo[1:]:
                propuestas.setdefault(valor, canonico)
        if campo == "destino" and conteo.get(VACIO):
            cercano = Counter(r.destino for r in renglones if r.destino).most_common(1)
            propuestas[VACIO] = cercano[0][0] if cercano else ""
        for valor, propuesto in propuestas.items():
            if valor != propuesto:
                salida.append((etiqueta, valor, conteo[valor], propuesto))
    return salida


def detectar_codigos(
    renglones: list[RenglonDiario], catalogo: dict[int, str]
) -> list[tuple[int, str, int, str]]:
    """(código, descripción en DIARIO, renglones, folios) de códigos fuera del catálogo."""
    grupos: dict[int, list[RenglonDiario]] = defaultdict(list)
    for r in renglones:
        if r.codigo is not None and r.codigo not in catalogo:
            grupos[r.codigo].append(r)
    salida = []
    for codigo, filas in sorted(grupos.items()):
        descripcion = Counter(r.descripcion for r in filas if r.descripcion).most_common(1)
        folios = sorted({r.folio for r in filas if r.folio is not None})
        salida.append(
            (
                codigo,
                descripcion[0][0] if descripcion else "",
                len(filas),
                ", ".join(map(str, folios)),
            )
        )
    return salida


def detectar_inventario(libro: LibroInventario, catalogo: dict[int, str]) -> list[tuple]:
    """(hoja, fila, código, descripción, dimensión, NP, cantidad, situación)."""
    salida = []
    por_np: dict[str, set[int]] = defaultdict(set)
    for hoja in libro.hojas:
        claves: dict[tuple, list] = defaultdict(list)
        for r in hoja.renglones:
            claves[(r.codigo, n.clave_estricta(r.dimension), n.clave_estricta(r.np))].append(r)
            for texto in (r.dimension, r.np):
                clave = n.clave_laxa(texto)
                if len(clave) >= 6:
                    por_np[clave].add(r.codigo)
            if r.codigo not in catalogo:
                salida.append(
                    (
                        hoja.nombre.strip(),
                        r.fila,
                        r.codigo,
                        r.descripcion,
                        r.dimension,
                        r.np,
                        r.cantidad,
                        "Código que no está en el catálogo ARTICULOS_MX",
                    )
                )
        for filas in claves.values():
            if len(filas) > 1:
                lista = ", ".join(str(x.fila) for x in filas)
                for r in filas:
                    salida.append(
                        (
                            hoja.nombre.strip(),
                            r.fila,
                            r.codigo,
                            r.descripcion,
                            r.dimension,
                            r.np,
                            r.cantidad,
                            f"Renglón repetido en la misma hoja (filas {lista}): ¿lotes distintos o doble captura?",
                        )
                    )
    for hoja in libro.hojas:
        for r in hoja.renglones:
            for texto in (r.dimension, r.np):
                codigos = por_np.get(n.clave_laxa(texto), set())
                if len(codigos) > 1:
                    salida.append(
                        (
                            hoja.nombre.strip(),
                            r.fila,
                            r.codigo,
                            r.descripcion,
                            r.dimension,
                            r.np,
                            r.cantidad,
                            f"'{texto}' aparece con los códigos {sorted(codigos)}: ¿posible doble conteo?",
                        )
                    )
                    break
    return salida


# ======================================================================
# Generación del Excel
# ======================================================================


def _hoja(
    libro: Workbook,
    titulo: str,
    encabezados: list[str],
    filas: list[list],
    anchos: list[int],
    editables: set[int],
    si_no: int | None = None,
    prioridad: int | None = None,
):
    ws = libro.create_sheet(titulo)
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(name=_FUENTE, bold=True, color="FFFFFF", size=10)
        c.fill = _ENCABEZADO
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = _BORDE
    for fila in filas:
        ws.append(fila)
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.font = Font(name=_FUENTE, size=10)
            c.border = _BORDE
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if c.column in editables:
                c.fill = _EDITABLE
        if prioridad and fila[prioridad - 1].value in _PRIORIDAD:
            fila[prioridad - 1].fill = _PRIORIDAD[fila[prioridad - 1].value]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{max(ws.max_row, 2)}"
    if si_no:
        validacion = DataValidation(type="list", formula1=f'"{SI},{NO}"', allow_blank=True)
        ws.add_data_validation(validacion)
        letra = get_column_letter(si_no)
        validacion.add(f"{letra}2:{letra}{max(ws.max_row, 2)}")
    return ws


def generar_revision(
    vales: LibroVales, inventario: LibroInventario, catalogo: dict[int, str], destino: Path
) -> Path:
    renglones = vales.renglones
    libro = Workbook()
    instrucciones = libro.active
    instrucciones.title = "Instrucciones"

    # 1 Renglones a corregir
    hallazgos = detectar_renglones(renglones, catalogo)
    sugerencias = sugerencias_encabezado(renglones)
    filas = []
    perdidas: list[tuple[int, int]] = []
    for i, h in enumerate(hallazgos, start=2):
        r = h.renglon
        fila = [h.prioridad, r.fila, r.folio, h.situacion]
        for columna, (_, campo) in enumerate(EDITABLES, start=len(FIJAS_HOJA1) + 1):
            valor = getattr(r, campo)
            if campo in r.errores or valor is None:
                valor = sugerencias.get(r.fila, {}).get(campo)
                if campo in r.errores:
                    perdidas.append((i, columna))
            if campo == "cantidad" and r.resto_cantidad:
                valor = r.crudo[9]
            fila.append(valor)
        fila += [SI if h.eliminar else None, None]
        filas.append(fila)
    encabezados = (
        FIJAS_HOJA1 + [e for e, _ in EDITABLES] + ["¿Eliminar renglón? (Sí/No)", "Comentario"]
    )
    editables = set(range(len(FIJAS_HOJA1) + 1, len(encabezados) + 1))
    ws = _hoja(
        libro,
        HOJA_RENGLONES,
        encabezados,
        filas,
        [8, 9, 7, 46] + [12] * len(EDITABLES) + [12, 24],
        editables,
        si_no=len(encabezados) - 1,
        prioridad=1,
    )
    for fila, columna in perdidas:
        ws.cell(fila, columna).fill = _PERDIDO
    ws.freeze_panes = "E2"
    columna_fecha = len(FIJAS_HOJA1) + 1
    for (celda,) in ws.iter_rows(min_row=2, min_col=columna_fecha, max_col=columna_fecha):
        celda.number_format = "dd/mm/yyyy"

    # 2 Duplicados
    duplicados = detectar_duplicados(renglones)
    filas = [
        [
            a.folio,
            a.fecha,
            a.fila,
            b.fila,
            a.cantidad,
            a.codigo,
            a.descripcion,
            a.clave,
            tipo,
            SI if tipo.startswith("Vale completo") else None,
            None,
        ]
        for a, b, tipo in duplicados
    ]
    _hoja(
        libro,
        HOJA_DUPLICADOS,
        [
            "Folio",
            "Fecha",
            "Fila original",
            "Fila repetida",
            "Cantidad",
            "Código",
            "Descripción",
            "Clave",
            "Tipo",
            "¿Eliminar la fila repetida? (Sí/No)",
            "Comentario",
        ],
        filas,
        [8, 11, 9, 9, 9, 8, 28, 20, 26, 16, 24],
        {10, 11},
        si_no=10,
    )

    # 3 Folios
    filas = [[p, f, s, q, None] for p, f, s, q in detectar_folios(renglones)]
    _hoja(
        libro,
        HOJA_FOLIOS,
        ["Prioridad", "Folio", "Situación", "Qué revisar", "Resultado (informativo)"],
        filas,
        [9, 10, 56, 48, 34],
        {5},
        prioridad=1,
    )

    # 4 Códigos
    filas = [[c, d, k, f, None, None, None] for c, d, k, f in detectar_codigos(renglones, catalogo)]
    _hoja(
        libro,
        HOJA_CODIGOS,
        [
            "Código en DIARIO",
            "Descripción en DIARIO",
            "Renglones",
            "Folios",
            "Código correcto (si es otro)",
            "Descripción correcta",
            "Comentario",
        ],
        filas,
        [10, 30, 9, 32, 14, 30, 24],
        {5, 6, 7},
    )

    # 5 Nombres
    filas = [[g, v, k, p, None] for g, v, k, p in detectar_nombres(renglones)]
    _hoja(
        libro,
        HOJA_NOMBRES,
        [
            "Grupo",
            "Variante encontrada",
            "Veces",
            "Nombre correcto (edítalo si no es)",
            "¿Es otra persona? (Sí/No)",
        ],
        filas,
        [8, 38, 8, 40, 14],
        {4, 5},
        si_no=5,
    )

    # 6 Valores
    filas = [[c, v, k, p, None] for c, v, k, p in detectar_valores(renglones)]
    _hoja(
        libro,
        HOJA_VALORES,
        [
            "Campo",
            "Valor encontrado",
            "Renglones",
            "Valor correcto (edítalo; vacío = dejar igual)",
            "Comentario",
        ],
        filas,
        [16, 30, 10, 40, 30],
        {4, 5},
    )

    # 7 Inventario
    filas = [list(x) + [None] for x in detectar_inventario(inventario, catalogo)]
    _hoja(
        libro,
        HOJA_INVENTARIO,
        [
            "Hoja",
            "Fila",
            "Código",
            "Descripción",
            "Dimensión",
            "NP",
            "Cantidad",
            "Situación",
            "Resultado (informativo)",
        ],
        filas,
        [28, 6, 8, 30, 22, 20, 9, 46, 30],
        {9},
    )

    _instrucciones(instrucciones, vales.ruta.name, inventario.ruta.name)
    libro.save(destino)
    return destino


def _instrucciones(ws, archivo_vales: str, archivo_inventario: str) -> None:
    ws["A1"] = "Revisión del historial antes de la primera carga"
    ws["A1"].font = Font(name=_FUENTE, bold=True, size=14)
    ws["A2"] = (
        f"Generado el {dt.date.today():%d/%m/%Y} a partir de: {archivo_vales} y {archivo_inventario}"
    )
    ws["A2"].font = Font(name=_FUENTE, italic=True, size=9, color="595959")
    ws["H1"] = MARCA_FORMATO
    ws["H1"].font = Font(name=_FUENTE, size=8, color="FFFFFF")
    lineas = [
        ("Cómo usar este archivo", True),
        (
            "1. Llena solo las celdas amarillas. Las naranjas son datos que se perdieron (#REF!): complétalos con el PDF escaneado.",
            False,
        ),
        (
            "2. Hoja 1: cada renglón ya trae los valores actuales o una sugerencia tomada de otros renglones del mismo folio. Corrige lo que no coincida con el PDF.",
            False,
        ),
        (
            "3. Hojas 2 y 5: responde Sí/No. Hojas 4 y 6: escribe el valor correcto solo si cambia.",
            False,
        ),
        (
            "4. Las hojas 3 y 7 son informativas: sirven para revisar físicamente o contra el PDF.",
            False,
        ),
        (
            "5. Guarda el archivo y cárgalo en la herramienta en el paso 'Primera carga'. Tus archivos originales no se modifican.",
            False,
        ),
        (
            "6. Puedes volver a generar esta lista cuantas veces quieras; la herramienta valida que cada fila siga correspondiendo al mismo folio.",
            False,
        ),
    ]
    for i, (texto, negrita) in enumerate(lineas, start=4):
        ws.cell(i, 1, texto).font = Font(name=_FUENTE, bold=negrita, size=11 if negrita else 10)
    ws.column_dimensions["A"].width = 120


# ======================================================================
# Lectura de respuestas
# ======================================================================


def _si(valor) -> bool:
    return (n.mayusculas(valor) or "").replace("Í", "I") in ("SI", "S")


def leer_revision(ruta: Path, renglones: list[RenglonDiario]) -> RespuestasRevision:
    libro = load_workbook(ruta, data_only=True)
    if (
        "Instrucciones" not in libro.sheetnames
        or libro["Instrucciones"]["H1"].value != MARCA_FORMATO
    ):
        raise ErrorRevision(
            "El archivo no es una lista de revisión generada por la herramienta (formato v2)."
        )
    por_fila = {r.fila: r for r in renglones}
    respuestas = RespuestasRevision()

    # 1 Renglones
    ws = libro[HOJA_RENGLONES]
    for valores in ws.iter_rows(min_row=2, values_only=True):
        if not valores or valores[1] is None:
            continue
        fila, folio = int(valores[1]), valores[2]
        actual = por_fila.get(fila)
        if actual is None or (
            folio is not None and actual.folio is not None and int(folio) != actual.folio
        ):
            respuestas.advertencias.append(
                f"Hoja 1: la fila {fila} ya no corresponde al folio {folio}; se ignoró."
            )
            continue
        if _si(valores[len(FIJAS_HOJA1) + len(EDITABLES)]):
            respuestas.eliminar.add(fila)
            continue
        cambios = {}
        for i, (_, campo) in enumerate(EDITABLES):
            nuevo = valores[len(FIJAS_HOJA1) + i]
            if campo == "fecha" and isinstance(nuevo, dt.datetime):
                nuevo = nuevo.date()
            elif campo == "fecha" and isinstance(nuevo, str):
                nuevo = n.fecha(nuevo)
            anterior = getattr(actual, campo)
            if campo == "cantidad":
                if n.decimal(nuevo) != anterior or actual.resto_cantidad:
                    cambios[campo] = nuevo
            elif campo in actual.errores or n.valor_a_texto(nuevo) != n.valor_a_texto(anterior):
                cambios[campo] = nuevo
        if cambios:
            respuestas.correcciones[fila] = cambios

    # 2 Duplicados
    for valores in libro[HOJA_DUPLICADOS].iter_rows(min_row=2, values_only=True):
        if valores and valores[3] is not None and _si(valores[9]):
            respuestas.eliminar.add(int(valores[3]))

    # 4 Códigos
    for valores in libro[HOJA_CODIGOS].iter_rows(min_row=2, values_only=True):
        if not valores or valores[0] is None:
            continue
        correcto = n.codigo_ax(valores[4])
        descripcion = n.valor_a_texto(valores[5])
        if correcto or descripcion:
            respuestas.codigos[int(valores[0])] = (correcto, descripcion)

    # 5 Nombres
    for valores in libro[HOJA_NOMBRES].iter_rows(min_row=2, values_only=True):
        if not valores or not valores[1]:
            continue
        variante, correcto = n.nombre_persona(valores[1]), n.nombre_persona(valores[3])
        if correcto and variante != correcto and not _si(valores[4]):
            respuestas.alias[variante] = correcto

    # 6 Valores
    for valores in libro[HOJA_VALORES].iter_rows(min_row=2, values_only=True):
        if not valores or not valores[0] or valores[1] is None:
            continue
        campo = CAMPOS_VALORES.get(valores[0])
        nuevo = n.valor_a_texto(valores[3])
        if campo and nuevo:
            respuestas.normalizaciones[(campo, str(valores[1]))] = None if nuevo == VACIO else nuevo
    return respuestas
