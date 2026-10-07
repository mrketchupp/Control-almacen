"""Genera Excel SINTÉTICOS que imitan la estructura de los archivos reales.

No contienen ningún dato real: nombres, códigos y cantidades son inventados.
Reproducen lo que importa para las pruebas:
- nombres de hoja con espacio final, tablas con fila de totales, fórmulas por renglón,
  notas de celda y hoja de catálogo oculta (inventario);
- hoja DIARIO con errores típicos (#REF!, duplicados, fechas como texto), hojas-formulario,
  macros, botón con macro asignada, imagen, customXml y calcChain (vales).
"""

from __future__ import annotations

import datetime as dt
import re
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

CATALOGO = {
    136: "SUMINISTRO DE DIESEL Y COMBUSTIBLE",
    701: "BALEROS",
    702: "FILTROS",
    703: "TORNILLOS",
    704: "EMPAQUE",
    705: "MANGUERA",
    706: "SELLOS",
    707: "VALVULA (INVENTARIABLE-REFACCIONES)",
    708: "GRASA",
    709: "CABLES",
    710: "LAMPARA DE USO GENERAL",
    711: "FOCO LED",
    714: "LAMPARA DE MANO",
}
SOLO_EN_VALES = {721: "CODIGO SOLO EN CATALOGO DE VALES"}

ENCABEZADOS_INV = [
    "ITEM",
    "CODIGO AX",
    "DESCRIPCIÓN ",
    "DIMENSION",
    "NP",
    "CANTIDAD",
    "UM",
    "CONSUMO",
    "INGRESO",
    "TOTAL",
]

# (item, código, dimensión, np, cantidad, um, consumo, ingreso)
HOJAS_INV = {
    "CONTENEDOR #1 INVENTARIABLE": [
        (1, 701, "6309-2Z/C3", None, 7, "PZA", None, None),
        (2, 701, "6205-2Z", None, 6, "PZA ", None, None),
        (3, 707, 'MARIPOSA 4"', None, 4, "PZA", None, None),
        (None, 710, "S/D", "X00489", 0, "PZA", None, None),
        (5, 706, 555001, None, 2, "PZA", 1, None),
    ],
    "CONTENEDOR #1 CONSUMIBLE ": [
        (1, 702, "P551317", None, 25, "PZA", None, None),
        (2, 702, "P557500", None, 21, "PZA", None, None),
        (3, 704, '6"', "FLEXITALIC", 65, "PZA", None, None),
        (4, 704, '6"', "FLEXITALIC", 65, "PZA", None, None),
        (5, 708, "ISOFLEX", None, 39, "PZA", None, 2),
        None,  # renglón vacío dentro de la tabla
    ],
    "CONTENEDOR #2 INVENTARIABLE": [
        (1, 701, "6309-2Z/C3", None, 3, "PZA", None, None),
        (2, 799, "SIN DIMENSION", None, 1, "PZA", None, None),
        # Ronda 15: en AX sin Tamaño ni Color. 711 tiene otra partida en AX con dimensión (MOD:A1): la
        # partida sin dimensión junta solo las variantes sin dimensión (S/D, SIN DIMENSION, S/N con NP
        # distintos). 714 solo tiene esa partida en AX: es el código completo, también MOD:HWD003.
        (3, 711, "S/D", "LED 20W", 4, "PZA", None, None),
        (4, 711, "SIN DIMENSION", "LED 50W", 3, "PZA", None, None),
        (5, 711, "S/N", "X100", 0, "PZA", None, None),
        (6, 711, "MOD:A1", None, 2, "PZA", None, None),
        (7, 714, "MOD:HWD003", None, 6, "PZA", None, None),
        (8, 714, "S/D", None, 2, "PZA", None, None),
    ],
    "CONTENEDOR #2 CONSUMIBLE": [
        (1, 705, '1/2"', None, 100, "MTS", None, None),
        (2, 705, "12", None, 5, "MTS", None, None),
    ],
}
NOTAS_INV = {("CONTENEDOR #1 INVENTARIABLE", "D4"): "Revisar dimensión en físico"}
PIE_DE_HOJA = {"CONTENEDOR #2 CONSUMIBLE": "TEXTO FUERA DE LA TABLA"}
# Encabezado de página como el del inventario real: título al centro y la fecha escrita a la derecha.
TITULO_INV = "INVENTARIO DE REFACCIONAMIENTO ALMACEN RIG 91 POZO PRUEBA 7"
FECHA_ENCABEZADO_INV = "LUNES 28 SEPTIEMBRE DE  2026"

ENCABEZADOS_DIARIO = [
    "FECHA",
    "No. folio",
    "Pase de Entrada",
    "Pase de Salida",
    "Origen: ",
    "Depto",
    "Destino",
    "Depto",
    "OC",
    "Cantidad",
    "Código",
    "Descripción",
    "CLAVE",
    "U.M.",
    "C.U",
    "Entrego/Recibio",
    "Entrego/Recibio",
    "Autorizo",
    "FAMILIA ",
    "TRANSFERENCIA/CONSUMO",
]
F = dt.datetime
_ALM, _MEC = "ALMACENISTA UNO", "MECANICO UNO"
# Cada renglón: 20 columnas A–T.
DIARIO = [
    [
        F(2026, 9, 1),
        1,
        0,
        "XXXXX",
        "RIG 91",
        "ALMACEN",
        "RIG 91",
        "MECANICO ",
        "S/OC",
        2,
        702,
        "FILTROS",
        "P551317",
        "PZA",
        0,
        _ALM,
        _MEC,
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 1),
        1,
        0,
        "XXXXX",
        "RIG 91",
        "ALMACEN",
        "RIG 91",
        "MECANICO ",
        "S/OC",
        3,
        702,
        "FILTROS",
        "P557500",
        "PZA",
        0,
        _ALM,
        _MEC,
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 2),
        2,
        0,
        "XXXXX",
        "RIG 91",
        "ALMACEN",
        "RIG 91",
        "OPERACIÓN",
        "S/OC",
        1,
        708,
        "GRASA",
        "ISOFLEX",
        "PZA",
        0,
        _ALM,
        "OPERADOR UNO",
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 2),
        2,
        0,
        "XXXXX",
        "RIG 91",
        "ALMACEN",
        "RIG 91",
        "OPERACIÓN",
        "S/OC",
        1,
        708,
        "GRASA",
        "ISOFLEX",
        "PZA",
        0,
        _ALM,
        "OPERADOR UNO",
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 2),
        3,
        "#REF!",
        "#REF!",
        "#REF!",
        "#REF!",
        "#REF!",
        "SOLDADOR",
        "S/OC",
        5,
        704,
        "EMPAQUE",
        '6"',
        "PZA",
        "#REF!",
        _ALM,
        "SOLDADOR UNO",
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 2),
        3,
        0,
        "XXXXX",
        "RIG 91",
        "MANTENIMIENTO",
        "RIG 91",
        "SOLDADOR",
        "S/OC",
        2,
        703,
        "TORNILLOS",
        "1/2 X 2",
        "PZA",
        0,
        _ALM,
        "SOLDADOR UNO",
        0,
        None,
        None,
    ],
    ["#REF!"] * 18 + [None, None],
    [
        " 03/09/2026",
        4,
        0,
        "XXXXX",
        "RIG 91",
        "ALMACEN",
        "RIG 91",
        "MECANICO ",
        "S/OC",
        "15LTS",
        799,
        "ARTICULO NUEVO",
        "S/D",
        "LTS",
        0,
        _ALM,
        _MEC,
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 3),
        5,
        0,
        "XXXXX",
        "RIG 91",
        "ALMACEN",
        0,
        "OPERACIÓN",
        "S/OC",
        200,
        136,
        "SUMINISTRO DE DIESEL Y COMBUSTIBLE",
        "DIESEL",
        "LTS",
        "MONTACARGAS",
        _ALM,
        "OPERADOR UNO",
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 4),
        6,
        0,
        "XXXXX",
        "RIG 91",
        "MANTENIMIENTO",
        "RIG 91",
        "TOP DRIVE ",
        "S/OC",
        1,
        701,
        "BALEROS",
        "6309-2Z/C3",
        "PZA",
        0,
        _ALM,
        "OPERADOR TD",
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 4),
        6,
        0,
        "XXXXX",
        "RIG 91",
        "MANTENIMIENTO",
        "RIG 91",
        "TOP DRIVE ",
        "S/OC",
        1,
        706,
        "SELLOS",
        555001,
        "PZA",
        0,
        _ALM,
        "OPERADOR TD",
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 5),
        7,
        0,
        "XXXXX",
        "RIG 91",
        "MANTENIMIENTO",
        "RIG 91",
        "MECANICO ",
        "S/OC",
        4,
        702,
        "FILTRO",
        "P551317",
        "PZA",
        0,
        _ALM,
        "MECANICO UNOO",
        0,
        None,
        None,
    ],
    [
        F(2026, 9, 5),
        7,
        0,
        "XXXXX",
        "RIG 91",
        "MANTENIMIENTO",
        "RIG 91",
        "MECANICO ",
        11536,
        1,
        704,
        "EMPAQUE",
        '6"',
        "PZA",
        "NUEVO",
        _ALM,
        "MECANICO UNOO",
        0,
        "INV",
        "CONSUMO",
    ],
    [
        F(2026, 9, 6),
        9,
        0,
        "XXXXX",
        "RIG 91 ",
        "ALMACEN",
        "RIG 48",
        "TRANSFERENCIA ",
        "S/CO",
        2,
        705,
        "MANGUERA",
        '1/2"',
        "MTS",
        0,
        _ALM,
        "RECEPTOR RIG 48",
        "AUTORIZADOR UNO",
        None,
        None,
    ],
]
FOLIO_CORTE = 5  # los vales con folio > 5 descuentan del conteo sintético

FORMULARIOS = {
    # hoja: (depto origen, depto destino, receptor, puesto receptor, desplazamiento firmas)
    "SOLDADOR": ("MANTENIMIENTO", "SOLDADOR", "SOLDADOR UNO", "SOLDADOR", 0),
    "MECANICO ": ("MANTENIMIENTO", "MECANICO ", _MEC, "MECANICO", -2),
    "TRANSFERENCIAS": ("ALMACEN", "TRANSFERENCIA ", "RECEPTOR RIG 48", "SUP. RIG 48", 0),
    # Como en el archivo real: en NOV el almacenista firma del lado derecho.
    "NOV": ("MANTENIMIENTO", "NOV ENERGY", "QUIMICO UNO", "QUIMICO", 0),
}
FIRMAS_INVERTIDAS = {"NOV"}
# Lugar de destino: los vales internos llegan al mismo equipo; NOV y transferencias no.
LUGAR_DESTINO = {"NOV": "RIG 91 - TANQUE NOV", "TRANSFERENCIAS": "RIG 48"}
# Renglones del formato por hoja (MECANICO tiene 19, como el real).
RENGLONES_FORMATO = {"MECANICO ": 19}

def png_color(ancho: int, alto: int, rgb: tuple[int, int, int]) -> bytes:
    """PNG de un solo color (fotos de ejemplo sintéticas)."""
    import struct
    import zlib

    def trozo(tipo: bytes, datos: bytes) -> bytes:
        return struct.pack(">I", len(datos)) + tipo + datos + struct.pack(">I", zlib.crc32(tipo + datos) & 0xFFFFFFFF)

    fila = b"\x00" + bytes(rgb) * ancho
    cabecera = struct.pack(">IIBBBBB", ancho, alto, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + trozo(b"IHDR", cabecera) + trozo(b"IDAT", zlib.compress(fila * alto)) + trozo(b"IEND", b"")


PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360000002000100e221bc330000000049454e44ae426082"
)


# ------------------------------------------------------------------ inventario


def generar_inventario(ruta: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for numero, (hoja, renglones) in enumerate(HOJAS_INV.items(), start=1):
        ws = wb.create_sheet(hoja)
        tabla = f"Tabla{numero}"
        ws.append(ENCABEZADOS_INV)
        for fila, datos in enumerate(renglones, start=2):
            if datos is None:
                ws.cell(
                    fila,
                    10,
                    f"={tabla}[[#This Row],[INGRESO]]+{tabla}[[#This Row],[CANTIDAD]]-{tabla}[[#This Row],[CONSUMO]]",
                )
                continue
            item, codigo, dimension, np, cantidad, um, consumo, ingreso = datos
            ws.cell(fila, 1, item)
            ws.cell(fila, 2, codigo)
            ws.cell(
                fila,
                3,
                f'=IF(B{fila}="","",VLOOKUP({tabla}[[#This Row],[CODIGO AX]],ARTICULOS_MX!$A$2:$B$5000,2,))',
            )
            ws.cell(fila, 4, dimension)
            ws.cell(fila, 5, np)
            ws.cell(fila, 6, cantidad)
            ws.cell(fila, 7, um)
            ws.cell(fila, 8, consumo)
            ws.cell(fila, 9, ingreso)
            ws.cell(
                fila,
                10,
                f"={tabla}[[#This Row],[INGRESO]]+{tabla}[[#This Row],[CANTIDAD]]-{tabla}[[#This Row],[CONSUMO]]",
            )
        totales = len(renglones) + 2
        ws.cell(totales, 1, "Total")
        for col, letra in ((6, "CANTIDAD"), (8, "CONSUMO"), (9, "INGRESO"), (10, "TOTAL")):
            ws.cell(totales, col, f"=SUBTOTAL(109,{tabla}[{letra}])")
        t = Table(
            displayName=tabla,
            name=tabla,
            ref=f"A1:J{totales}",
            totalsRowCount=1,
            autoFilter=AutoFilter(ref=f"A1:J{totales - 1}"),
            tableStyleInfo=TableStyleInfo(name="TableStyleMedium2", showRowStripes=True),
        )
        t.tableColumns = [
            TableColumn(
                id=i,
                name=nombre,
                totalsRowLabel="Total" if i == 1 else None,
                totalsRowFunction="sum"
                if nombre in ("CANTIDAD", "CONSUMO", "INGRESO", "TOTAL")
                else None,
            )
            for i, nombre in enumerate(ENCABEZADOS_INV, start=1)
        ]
        ws.add_table(t)
        ws.print_area = f"A1:J{totales}"
        if hoja in PIE_DE_HOJA:
            ws.cell(totales + 2, 1, PIE_DE_HOJA[hoja])
        ws.oddHeader.center.text = f"\n{TITULO_INV}\n&A"
        ws.oddHeader.center.font = "-,Negrita"
        ws.oddHeader.center.size = 16
        ws.oddHeader.right.text = f"{FECHA_ENCABEZADO_INV}\n"
        ws.oddFooter.center.text = "ALMACENISTA: FULANO DE TAL"
        for (h, celda), texto in NOTAS_INV.items():
            if h == hoja:
                ws[celda].comment = Comment(f"AUTOR:\n{texto}", "AUTOR")
    cat = wb.create_sheet("ARTICULOS_MX")
    cat.append([None, None])
    cat.append(["CODIGO AX", "PRODUCTO"])
    for codigo, descripcion in CATALOGO.items():
        cat.append([codigo, descripcion])
    cat.sheet_state = "hidden"
    wb.save(ruta)
    _agregar_calcchain(ruta, [("C2", 1), ("J2", 1)])
    return ruta


# ---------------------------------------------------------------------- vales


def generar_vales(ruta: Path) -> Path:
    wb = Workbook()
    diario = wb.active
    diario.title = "DIARIO"
    diario.append(ENCABEZADOS_DIARIO)
    for renglon in DIARIO:
        diario.append(renglon)
    for fila in range(2, len(DIARIO) + 2):
        diario.cell(fila, 1).number_format = "dd/mm/yyyy"
    diario.freeze_panes = "A2"
    diario.auto_filter.ref = f"A1:T{len(DIARIO) + 1}"

    delgado = Side(style="thin", color="000000")
    borde = Border(left=delgado, right=delgado, top=delgado, bottom=delgado)
    for hoja, (d_origen, d_destino, receptor, puesto, dz) in FORMULARIOS.items():
        ws = wb.create_sheet(hoja)
        renglones = RENGLONES_FORMATO.get(hoja, 21)
        # Formato visual como el real: anchos, celdas combinadas, colores, bordes y página.
        for col, ancho in zip("ABCDEFGHIJKL", [4, 11.4, 11.7, 13.7, 22.1, 12.4, 11.9, 31.4, 26.6, 14.1, 21.1, 12.1], strict=True):
            ws.column_dimensions[col].width = ancho
        for rango in ("J6:K6", "C17:D17", "E17:G17", "I17:K17", "C18:D18", "E18:G18", "I18:K18", "F20:H20"):
            ws.merge_cells(rango)
        for fila in range(21, 21 + renglones):
            ws.merge_cells(f"F{fila}:H{fila}")
            for col in "CDEFGHIJK":
                ws[f"{col}{fila}"].border = borde
        for fila in range(21 + renglones, 42):
            ws.merge_cells(f"C{fila}:K{fila}")
        for col in "CDEFGHIJK":
            celda = ws[f"{col}20"]
            celda.font = Font(name="Arial", bold=True, size=10)
            celda.fill = PatternFill("solid", fgColor="FFFF00" if col != "F" else "E26B0A")
            celda.alignment = Alignment(horizontal="center")
            celda.border = borde
        ws["C17"].font = ws["C18"].font = Font(name="Arial", bold=True)
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = "portrait"
        ws.page_setup.scale = 59
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.print_options.horizontalCentered = True
        ws.page_margins.left, ws.page_margins.right = 0.59, 0.39
        ws.page_margins.top, ws.page_margins.bottom = 0.79, 0.59
        ws.oddFooter.left.text = "FORMATO-PRUEBA"
        ws.oddFooter.right.text = "Emision: X"
        ws["I6"], ws["J6"] = "Fecha:", "=TODAY()"
        ws["J8"], ws["K8"] = "No. folio", 9
        ws["J10"], ws["J11"], ws["K11"] = "Entradas", "Salida Planta", "XXXXX"
        ws["C17"], ws["E17"], ws["H17"], ws["I17"] = (
            "Origen: ",
            "RIG 91",
            "Departamento:  ",
            d_origen,
        )
        ws["C18"], ws["E18"], ws["H18"], ws["I18"] = (
            "Destino: ",
            LUGAR_DESTINO.get(hoja, "RIG 91"),
            "Departamento: ",
            d_destino,
        )
        for col, texto in zip(
            "CDEFIJK",
            [
                "O.C.",
                "CANTIDAD",
                "CODIGO",
                "DESCRIPCION DEL MATERIAL",
                "CLAVE ALMACEN",
                "PRESENTACION",
                "LOTE",
            ],
            strict=True,
        ):
            ws[f"{col}20"] = texto
        for fila in range(21, 21 + renglones):
            ws[f"F{fila}"] = f'=IF(E{fila}="","",VLOOKUP(E{fila},$AG$6:$AH$10030,2,0))'
        ws[f"C{42 + dz}"] = "OBSERVACION"
        ws[f"C{44 + dz}"] = (
            "ESTE MATERIAL CUMPLE CON LAS ESPECIFICACIONES REQUERIDAS POR EL USUARIO"
        )
        ws[f"C{45 + dz}"] = "MATERIAL SUMINISTRADO PARA USO EN MANTENIMIENTO DEL RIG 91"
        ws[f"C{46 + dz}"] = 'ETAPA DE PERFORACION: 8 1/2"'
        # Fila espaciadora de 1 px con un número olvidado: Excel no lo muestra.
        ws[f"C{47 + dz}"] = 3
        ws.row_dimensions[47 + dz].height = 0.75
        izquierda = (receptor, puesto) if hoja in FIRMAS_INVERTIDAS else (_ALM, "ALMACENISTA")
        derecha = (_ALM, "ALMACENISTA") if hoja in FIRMAS_INVERTIDAS else (receptor, puesto)
        if hoja in FIRMAS_INVERTIDAS:
            ws[f"D{50 + dz}"], ws[f"J{50 + dz}"] = "RECIBE / AUTORIZA ", "ENTREGA / AUTORIZA "
        else:
            ws[f"D{50 + dz}"], ws[f"J{50 + dz}"] = "ENTREGO/RECIBIO", "RECIBIO/ENTREGO"
        ws.merge_cells(f"D{52 + dz}:F{52 + dz}")
        ws.merge_cells(f"I{52 + dz}:K{52 + dz}")
        ws[f"C{52 + dz}"], ws[f"D{52 + dz}"], ws[f"I{52 + dz}"] = "Nombre:", izquierda[0], derecha[0]
        ws[f"C{53 + dz}"], ws[f"D{53 + dz}"], ws[f"I{53 + dz}"] = "Puesto:", izquierda[1], derecha[1]
        if hoja == "TRANSFERENCIAS":
            ws["G56"], ws["F57"] = "AUTORIZA", "Firma:"
            ws["F58"], ws["G58"] = "Nombre:", "AUTORIZADOR UNO"
            ws["F59"], ws["G59"] = "Puesto:", "RIG MANAGER"
        if hoja in FIRMAS_INVERTIDAS:
            # Como NOV real: segunda fila de firmas (personal de la compañía y patrimonial).
            ws.merge_cells("D56:F56")
            ws.merge_cells("I56:K56")
            ws["C55"], ws["E55"], ws["J55"] = "FIRMA:", "RECIBE / AUTORIZA ", "PATRIMONIAL"
            ws["C56"], ws["D56"], ws["I56"] = "NOMBRE:", "PERSONAL NOV UNO", "PATRIMONIAL UNO"
            ws["C57"], ws["D57"], ws["I57"] = "PUESTO:", "NOV ENERGY", "SEG PATRIMONIAL"
        ws["AG5"], ws["AH5"] = "CODIGO AX", "PRODUCTO"
        for fila, (codigo, descripcion) in enumerate(
            {**CATALOGO, **SOLO_EN_VALES}.items(), start=6
        ):
            ws.cell(fila, 33, codigo)
            ws.cell(fila, 34, descripcion)
        ws.print_area = "C5:K62"
    wb.save(ruta)
    _convertir_en_xlsm(ruta)
    return ruta


# --------------------------------------------------------- post-proceso de ZIP


def _reescribir_zip(ruta: Path, cambios: dict[str, bytes | None]) -> None:
    with zipfile.ZipFile(ruta) as origen:
        partes = [(info, origen.read(info.filename)) for info in origen.infolist()]
    nombres = {info.filename for info, _ in partes}
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as destino:
        for info, datos in partes:
            if info.filename in cambios:
                nuevo = cambios[info.filename]
                if nuevo is not None:
                    destino.writestr(info, nuevo)
            else:
                destino.writestr(info, datos)
        for nombre, datos in cambios.items():
            if nombre not in nombres and datos is not None:
                destino.writestr(nombre, datos)


def _leer(ruta: Path, parte: str) -> str:
    with zipfile.ZipFile(ruta) as z:
        return z.read(parte).decode("utf-8")


def _agregar_calcchain(ruta: Path, celdas: list[tuple[str, int]]) -> None:
    tipos = _leer(ruta, "[Content_Types].xml").replace(
        "</Types>",
        '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>',
    )
    rels = _leer(ruta, "xl/_rels/workbook.xml.rels").replace(
        "</Relationships>",
        '<Relationship Id="rIdCalc" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>',
    )
    cadena = "".join(f'<c r="{r}" i="{i}"/>' for r, i in celdas)
    calc = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">{cadena}</calcChain>'
    _reescribir_zip(
        ruta,
        {
            "[Content_Types].xml": tipos.encode(),
            "xl/_rels/workbook.xml.rels": rels.encode(),
            "xl/calcChain.xml": calc.encode(),
        },
    )


def _convertir_en_xlsm(ruta: Path) -> None:
    """Agrega macros, un botón con macro, una imagen y customXml (como el archivo real)."""
    tipos = _leer(ruta, "[Content_Types].xml")
    tipos = tipos.replace(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    ).replace(
        "</Types>",
        '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/>'
        '<Default Extension="png" ContentType="image/png"/>'
        '<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>'
        '<Override PartName="/customXml/itemProps1.xml" ContentType="application/vnd.openxmlformats-officedocument.customXmlProperties+xml"/>'
        "</Types>",
    )
    rels = _leer(ruta, "xl/_rels/workbook.xml.rels").replace(
        "</Relationships>",
        '<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
        '<Relationship Id="rIdCx1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml" Target="../customXml/item1.xml"/>'
        "</Relationships>",
    )
    hoja2 = _leer(ruta, "xl/worksheets/sheet2.xml")
    hoja2 = re.sub(r"(<pageMargins[^>]*/>)", r'\1<drawing r:id="rIdDib1"/>', hoja2, count=1)
    if 'xmlns:r="' not in hoja2:
        hoja2 = hoja2.replace(
            "<worksheet ",
            '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ',
            1,
        )
    try:
        rels_hoja2 = _leer(ruta, "xl/worksheets/_rels/sheet2.xml.rels")
    except KeyError:
        rels_hoja2 = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>'
    rels_hoja2 = rels_hoja2.replace(
        "</Relationships>",
        '<Relationship Id="rIdDib1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing1.xml"/></Relationships>',
    )
    dibujo = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        "<xdr:twoCellAnchor><xdr:from><xdr:col>11</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>9</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
        "<xdr:to><xdr:col>13</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>11</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>"
        '<xdr:sp macro="[0]!PasarDatos" textlink=""><xdr:nvSpPr><xdr:cNvPr id="2" name="Boton GRABAR"/><xdr:cNvSpPr/></xdr:nvSpPr>'
        '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr>'
        "<xdr:txBody><a:bodyPr/><a:p><a:r><a:t>GRABAR </a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>"
        "<xdr:oneCellAnchor><xdr:from><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>5</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
        '<xdr:ext cx="952500" cy="317500"/><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="3" name="Logo"/><xdr:cNvPicPr/></xdr:nvPicPr>'
        '<xdr:blipFill><a:blip r:embed="rIdImg1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
        '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor>'
        "</xdr:wsDr>"
    )
    rels_dibujo = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdImg1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/>'
        "</Relationships>"
    )
    def foto(num: int, col1: int, col2: int, rid: str) -> str:
        return (
            f"<xdr:twoCellAnchor><xdr:from><xdr:col>{col1}</xdr:col><xdr:colOff>95250</xdr:colOff><xdr:row>24</xdr:row><xdr:rowOff>57150</xdr:rowOff></xdr:from>"
            f"<xdr:to><xdr:col>{col2}</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>37</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>"
            f'<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{10 + num}" name="Foto {num}"/><xdr:cNvPicPr/></xdr:nvPicPr>'
            f'<xdr:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
            '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:twoCellAnchor>'
        )

    dibujo_nov = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        "<xdr:oneCellAnchor><xdr:from><xdr:col>2</xdr:col><xdr:colOff>0</xdr:colOff><xdr:row>4</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from>"
        '<xdr:ext cx="952500" cy="317500"/><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="3" name="Logo"/><xdr:cNvPicPr/></xdr:nvPicPr>'
        '<xdr:blipFill><a:blip r:embed="rIdImg1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
        '<xdr:spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic><xdr:clientData/></xdr:oneCellAnchor>'
        + foto(1, 3, 4, "rIdFoto1")
        + foto(2, 5, 8, "rIdFoto2")
        + foto(3, 9, 10, "rIdFoto3")
        + "</xdr:wsDr>"
    )
    rels_dibujo_nov = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdImg1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image1.png"/>'
        '<Relationship Id="rIdFoto1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/foto1.png"/>'
        '<Relationship Id="rIdFoto2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/foto2.png"/>'
        '<Relationship Id="rIdFoto3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/foto3.png"/>'
        "</Relationships>"
    )
    hoja5 = _leer(ruta, "xl/worksheets/sheet5.xml")
    hoja5 = re.sub(r"(<pageMargins[^>]*/>)", r'\1<drawing r:id="rIdDib2"/>', hoja5, count=1)
    if 'xmlns:r="' not in hoja5:
        hoja5 = hoja5.replace(
            "<worksheet ",
            '<worksheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" ',
            1,
        )
    rels_hoja5 = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rIdDib2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="../drawings/drawing2.xml"/></Relationships>'
    )
    tipos = tipos.replace(
        "</Types>",
        '<Override PartName="/xl/drawings/drawing2.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>',
    )
    custom = '<?xml version="1.0" encoding="UTF-8" standalone="no"?><p:properties xmlns:p="http://schemas.microsoft.com/office/2006/metadata/properties"/>'
    props = (
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
        '<ds:datastoreItem ds:itemID="{00000000-0000-0000-0000-000000000001}" '
        'xmlns:ds="http://schemas.openxmlformats.org/officeDocument/2006/customXml"/>'
    )
    rels_custom = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps" Target="itemProps1.xml"/>'
        "</Relationships>"
    )
    _reescribir_zip(
        ruta,
        {
            "[Content_Types].xml": tipos.encode(),
            "xl/_rels/workbook.xml.rels": rels.encode(),
            "xl/worksheets/sheet2.xml": hoja2.encode(),
            "xl/worksheets/_rels/sheet2.xml.rels": rels_hoja2.encode(),
            "xl/drawings/drawing1.xml": dibujo.encode(),
            "xl/drawings/_rels/drawing1.xml.rels": rels_dibujo.encode(),
            "xl/media/image1.png": PNG_1PX,
            "xl/worksheets/sheet5.xml": hoja5.encode(),
            "xl/worksheets/_rels/sheet5.xml.rels": rels_hoja5.encode(),
            "xl/drawings/drawing2.xml": dibujo_nov.encode(),
            "xl/drawings/_rels/drawing2.xml.rels": rels_dibujo_nov.encode(),
            "xl/media/foto1.png": png_color(30, 36, (200, 170, 90)),
            "xl/media/foto2.png": png_color(40, 30, (40, 90, 160)),
            "xl/media/foto3.png": png_color(30, 34, (90, 140, 70)),
            "xl/vbaProject.bin": b"VBA-FICTICIO-SOLO-PARA-PRUEBAS" * 64,
            "customXml/item1.xml": custom.encode(),
            "customXml/itemProps1.xml": props.encode(),
            "customXml/_rels/item1.xml.rels": rels_custom.encode(),
        },
    )
    _agregar_calcchain(ruta, [("F21", 2)])


# ---------------------------------------------------------------- reporte AX (conciliación, F4)
# Inventario de AX al 05-09-2026 (los vales hasta el 05 ya están en AX; el folio 9, del 06, va en tránsito).
# Casos: exactos, Tamaño cortado a 10 caracteres (MARIPOSA 4), NP en Color (FLEXITALIC), error de dedo
# (P55I317), '1/2"' contra '12', renglones de AX sin físico (703, 709), otro almacén que se filtra,
# diferencia explicada por un vale en tránsito (705 1/2") y diferencias sin explicar.
HOJA_AX = "rptInventSumDateTransForDimensi"
ENCABEZADOS_AX = [
    "Código de Artículo",
    "Nombre del Artículo",
    "Modelo de Inventario",
    "Unidad de Medida",
    "Almacén",
    "Tamaño",
    "Color",
    "Disponible",
    "Valor Financiero",
    "Valor de Inventario",
]
ALMACEN_AX = "RIG91-IX25"
# (código, tamaño, color, um, disponible, costo unitario, almacén)
RENGLONES_AX = [
    (701, "6309-2Z/C3", "", "PZA", 9, 350.0, ALMACEN_AX),
    (701, "6205-2Z", "", "PZA", 6, 120.5, ALMACEN_AX),
    (702, "P55I317", "", "PZA", 21, 85.25, ALMACEN_AX),
    (702, "P557500", "", "PZA", 20, 85.25, ALMACEN_AX),
    (704, '6"', "FLEXITALIC", "PZA", 65, 40.0, ALMACEN_AX),
    (705, '1/2"', "", "m", 100, 12.5, ALMACEN_AX),
    (705, "12", "", "m", 5, 30.0, ALMACEN_AX),
    (706, "555001", "", "PZA", 1, 900.0, ALMACEN_AX),
    (707, "MARIPOSA 4", "", "PZA", 4, 1500.0, ALMACEN_AX),
    (708, "ISOFLEX", "", "PZA", 37, 210.0, ALMACEN_AX),
    (710, "S/D", "X00489", "PZA", 2, 75.0, ALMACEN_AX),
    (703, "1/2 X 2", "", "PZA", 10, 4.75, ALMACEN_AX),
    (709, "CABLE 3/4", "", "m", 50, 18.0, ALMACEN_AX),
    (711, "", "", "PZA", 7, 60.0, ALMACEN_AX),
    (711, "MOD:A1", "", "PZA", 2, 60.0, ALMACEN_AX),
    (714, "", "", "PZA", 8, 45.0, ALMACEN_AX),
    (701, "6309-2Z/C3", "", "PZA", 99, 350.0, "RIG48-XX10"),
    # Otro modelo de inventario (no INV, como el diésel): no se concilia.
    (136, "", "", "LT", 500, 20.0, ALMACEN_AX, "NO INV"),
]
NOMBRE_AX = "DELTA RIG 91 SINTETICO 050926.xlsx"


def generar_ax(ruta: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA_AX
    ws.append(ENCABEZADOS_AX)
    for codigo, tamano, color, um, disponible, costo, almacen, *modelo in RENGLONES_AX:
        valor = round(disponible * costo, 2)
        ws.append([f"{codigo:09d}", CATALOGO.get(codigo, "ARTICULO"), modelo[0] if modelo else "INV", um, almacen, tamano, color, disponible, valor, valor])
    for fila in range(2, ws.max_row + 1):
        ws.cell(fila, 8).number_format = "#,##0.00"
        ws.cell(fila, 9).number_format = "#,##0.00"
        ws.cell(fila, 10).number_format = "#,##0.00"
    wb.save(ruta)
    return ruta


# ---------------------------------------------------------------- archivo de la base

# El encargado de la base copia el DIARIO de salidas y le agrega estas columnas (S–X):
# INV/NINV (si se descuenta en AX), TIPO DE MOV (consumo o transferencia), CANTIDAD (lo aplicado),
# TR / IN (folio que da AX) y COMENTARIOS. Solo cambian S y T respecto al DIARIO de la herramienta.
ENCABEZADOS_BASE = ENCABEZADOS_DIARIO[:18] + ["INV/NINV", "TIPO DE MOV", "CANTIDAD ", "TR", "IN", "COMENTARIOS"]
NOMBRE_BASE = "VALES DE SALIDA BASE SINTETICO 050926.xlsx"
# (folio, código, clave, cantidad, INV/NINV, TIPO DE MOV, CANTIDAD aplicada, TR, IN, COMENTARIOS)
SEGUIMIENTO_BASE = [
    (1, 702, "P551317", 2, "INV", "CONSUMO", 2, None, "IN00000101", None),  # aplicada
    (1, 702, "P557500", 3, "INV", "CONSUMO", 1, None, "IN00000101", "FALTAN 2"),  # parcial: 2 pendientes
    (2, 708, "ISOFLEX", 1, "INV", "TRANSFERENCIA", None, "TRS000000201", None, None),  # U vacía = todo
    (3, 704, '6"', 5, "NO INV", None, None, None, None, None),  # no se descuenta en AX
    (3, 703, "1/2 X 2", 2, "CONPROV", None, None, None, None, None),
    (4, 799, "S/D", 15, None, None, None, None, None, None),  # sin revisar
    (5, 136, "DIESEL", 200, "NO INV ", None, None, None, None, None),
    (6, 701, "6309-2Z/C3", 1, "INV", "CONSUMO", None, "PENDIENTE", None, None),  # pendiente
    (6, 706, "555001", 1, "INV", None, None, None, None, None),  # pendiente
    (7, 702, "P551318", 4, "INV", "CONSUMO", 4, None, "IN00000150", None),  # clave distinta
    (7, 704, '6"', 1, "INV", "TRANSFERENCIA", "REGRESAR", "TRS000000233", None, None),  # U no numérica
    (7, 712, "6303 SKF", 1, "INV", "CONSUMO", 1, None, "IN00000150", None),  # no está en la herramienta
]


def generar_base(ruta: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "DIARIO"
    ws.append(ENCABEZADOS_BASE)
    fechas = {r[1]: r[0] for r in DIARIO if isinstance(r[1], int) and isinstance(r[0], dt.datetime)}
    for folio, codigo, clave, cantidad, inv, mov, aplicada, tr, in_, comentario in SEGUIMIENTO_BASE:
        base = [fechas.get(folio), folio, 0, "XXXXX", "RIG 91", "ALMACEN", "RIG 91", "MECANICO ", "S/OC", cantidad, codigo,
                CATALOGO.get(codigo, "ARTICULO"), clave, "PZA", 0, "ALMACENISTA UNO", "MECANICO UNO", 0]
        ws.append(base + [inv, mov, aplicada, tr, in_, comentario])
    oculta = wb.create_sheet("Produccion Nueva 2010")
    oculta.sheet_state = "hidden"
    oculta.append(["NO ES EL DIARIO"])
    wb.save(ruta)
    return ruta


def generar_todo(carpeta: Path) -> dict[str, Path]:
    carpeta.mkdir(parents=True, exist_ok=True)
    return {
        "inventario": generar_inventario(carpeta / "INVENTARIO SINTETICO.xlsx"),
        "vales": generar_vales(carpeta / "VALES SINTETICO.xlsm"),
        "ax": generar_ax(carpeta / NOMBRE_AX),
        "base": generar_base(carpeta / NOMBRE_BASE),
    }


if __name__ == "__main__":
    # Uso: python tests/fixtures/generar.py <carpeta>
    # Imprime en JSON las rutas generadas y las constantes que usan las pruebas (tests/ayuda.js).
    import json
    import sys

    destino = Path(sys.argv[1]) if len(sys.argv) > 1 else Path.cwd()
    rutas = generar_todo(destino)
    print(
        json.dumps(
            {
                **{nombre: str(ruta) for nombre, ruta in rutas.items()},
                "catalogo": {str(k): v for k, v in CATALOGO.items()},
                "folio_corte": FOLIO_CORTE,
            },
            ensure_ascii=False,
        )
    )
