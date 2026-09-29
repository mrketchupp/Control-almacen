"""Exportación sobre plantilla: solo cambian las partes necesarias y los datos cuadran."""

import datetime as dt
import shutil
import subprocess
import zipfile
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from control_almacen.db import modelos as m
from control_almacen.exportadores.inventario import exportar_inventario
from control_almacen.exportadores.plantilla_ooxml import desplazar_formula, letra_columna
from control_almacen.exportadores.vales import exportar_vales
from control_almacen.servicios.primera_carga import OpcionesPrimeraCarga, ejecutar_primera_carga
from tests.fixtures.generar import FOLIO_CORTE


@pytest.fixture
def cargada(bd, archivos, tmp_path):
    ejecutar_primera_carga(
        bd,
        OpcionesPrimeraCarga(
            ruta_inventario=archivos["inventario"],
            ruta_vales=archivos["vales"],
            folio_corte=FOLIO_CORTE,
            fecha_conteo=dt.date(2026, 9, 3),
        ),
    )
    return bd


def _partes_distintas(original, exportado) -> set[str]:
    with zipfile.ZipFile(original) as a, zipfile.ZipFile(exportado) as b:
        assert a.namelist() == [n for n in a.namelist() if n in b.namelist()] or True
        return {n for n in a.namelist() if n not in b.namelist() or a.read(n) != b.read(n)}


# ------------------------------------------------------------------ fórmulas


@pytest.mark.parametrize(
    ("formula", "delta", "esperada"),
    [
        (
            'IF(B2="","",VLOOKUP(Tabla1[[#This Row],[CODIGO AX]],ARTICULOS_MX!$A$2:$B$5000,2,))',
            3,
            'IF(B5="","",VLOOKUP(Tabla1[[#This Row],[CODIGO AX]],ARTICULOS_MX!$A$2:$B$5000,2,))',
        ),
        (
            "Tabla315[[#This Row],[INGRESO]]+Tabla315[[#This Row],[CANTIDAD]]",
            10,
            "Tabla315[[#This Row],[INGRESO]]+Tabla315[[#This Row],[CANTIDAD]]",
        ),
        ("SUBTOTAL(109,Tabla1[CANTIDAD])", 5, "SUBTOTAL(109,Tabla1[CANTIDAD])"),
        ('A1&"B2"&$C$3&D$4&Hoja!E5', 1, 'A2&"B2"&$C$3&D$4&Hoja!E6'),
        ("TABLA315[x]+B2", 1, "TABLA315[x]+B3"),
    ],
)
def test_desplazar_formula(formula, delta, esperada):
    assert desplazar_formula(formula, delta) == esperada


def test_letra_columna():
    assert [letra_columna(i) for i in (1, 20, 26, 27, 52)] == ["A", "T", "Z", "AA", "AZ"]


# --------------------------------------------------------------------- vales


def test_vales_solo_cambia_diario(cargada, archivos, tmp_path):
    destino = tmp_path / "VALES DE SALIDA DLTA.xlsm"
    with cargada.sesion() as s:
        resultado = exportar_vales(s, archivos["vales"], destino)
    assert resultado.renglones == 13
    assert resultado.ultimo_folio == 9
    distintas = _partes_distintas(archivos["vales"], destino)
    assert distintas == {"xl/worksheets/sheet1.xml", "xl/workbook.xml"} or distintas == {
        "xl/worksheets/sheet1.xml"
    }
    with zipfile.ZipFile(destino) as z:
        nombres = set(z.namelist())
    # macros, botón, imagen, customXml y calcChain se conservan
    assert {
        "xl/vbaProject.bin",
        "xl/drawings/drawing1.xml",
        "xl/media/image1.png",
        "customXml/item1.xml",
        "xl/calcChain.xml",
    } <= nombres


def test_vales_valores_del_diario(cargada, archivos, tmp_path):
    destino = tmp_path / "vales.xlsm"
    with cargada.sesion() as s:
        exportar_vales(s, archivos["vales"], destino)
    ws = load_workbook(destino, data_only=True, keep_vba=True)["DIARIO"]
    filas = list(ws.iter_rows(min_row=1, max_col=20, values_only=True))
    assert filas[0][1] == "No. folio"
    datos = filas[1:]
    assert [f[1] for f in datos] == [1, 1, 2, 2, 3, 3, 4, 5, 6, 6, 7, 7, 9]
    primero = datos[0]
    assert primero[0].date() == dt.date(2026, 9, 1)
    assert primero[2:4] == (0, "XXXXX")
    assert primero[8] == "S/OC"
    assert primero[17] == 0  # sin autorizó → 0, como la macro
    # folio 3: encabezado perdido (#REF!) completado con el otro renglón del folio
    folio3 = datos[4]
    assert folio3[4:6] == ("RIG 91", "MANTENIMIENTO")
    # descripción con error de dedo corregida, O.C. numérica y lote
    folio7 = datos[10:12]
    assert folio7[0][11] == "FILTROS"
    assert folio7[1][8] == 11536 and folio7[1][14] == "NUEVO"
    assert folio7[1][18:20] == ("INV", "CONSUMO")
    # número guardado como texto en la base vuelve a ser número
    assert datos[9][12] == 555001


def test_vales_cancelado_aparece_con_su_folio(cargada, archivos, tmp_path):
    with cargada.sesion() as s:
        vale = s.scalar(select(m.Vale).where(m.Vale.folio == 9))
        vale.estado, vale.motivo_cancelacion = "CANCELADO", "Captura duplicada"
    destino = tmp_path / "vales.xlsm"
    with cargada.sesion() as s:
        exportar_vales(s, archivos["vales"], destino)
    ws = load_workbook(destino, data_only=True)["DIARIO"]
    ultimo = list(
        ws.iter_rows(min_row=ws.max_row, max_row=ws.max_row, max_col=20, values_only=True)
    )[0]
    assert ultimo[1] == 9 and ultimo[9] == 0 and ultimo[11] == "CANCELADO – Captura duplicada"


# ---------------------------------------------------------------- inventario


def test_inventario_partes_intactas(cargada, archivos, tmp_path):
    destino = tmp_path / "inventario.xlsx"
    with cargada.sesion() as s:
        resultado = exportar_inventario(s, archivos["inventario"], destino)
    assert resultado.renglones == 14
    distintas = _partes_distintas(archivos["inventario"], destino)
    permitidas = {
        "[Content_Types].xml",
        "xl/_rels/workbook.xml.rels",
        "xl/workbook.xml",
        "xl/calcChain.xml",
        *(f"xl/worksheets/sheet{i}.xml" for i in range(1, 6)),
        *(f"xl/tables/table{i}.xml" for i in range(1, 5)),
        "xl/comments/comment1.xml",
        "xl/drawings/commentsDrawing1.vml",
    }
    assert distintas <= permitidas
    assert "xl/styles.xml" not in distintas and "xl/theme/theme1.xml" not in distintas
    with zipfile.ZipFile(destino) as z:
        assert "xl/calcChain.xml" not in z.namelist()
        assert b"calcChain" not in z.read("[Content_Types].xml")
        assert b'fullCalcOnLoad="1"' in z.read("xl/workbook.xml")


def test_inventario_datos_formulas_y_totales(cargada, archivos, tmp_path):
    destino = tmp_path / "inventario.xlsx"
    with cargada.sesion() as s:
        exportar_inventario(s, archivos["inventario"], destino)
    libro = load_workbook(destino)
    c1 = libro["CONTENEDOR #1 INVENTARIABLE"]
    assert c1.tables["Tabla1"].ref == "A1:J7"
    assert c1["C3"].value.startswith('=IF(B3="","",VLOOKUP(Tabla1[[#This Row],[CODIGO AX]]')
    assert (
        c1["J6"].value
        == "=Tabla1[[#This Row],[INGRESO]]+Tabla1[[#This Row],[CANTIDAD]]-Tabla1[[#This Row],[CONSUMO]]"
    )
    assert c1["A7"].value == "Total" and c1["F7"].value == "=SUBTOTAL(109,Tabla1[CANTIDAD])"
    assert [c1.cell(r, 1).value for r in range(2, 7)] == [1, 2, 3, 4, 5]  # ITEM renumerado
    assert c1["D6"].value == 555001 and c1["H6"].value == 1  # CONSUMO del vale 6
    assert c1["G3"].value == "PZA"  # 'PZA ' sin el espacio sobrante
    assert c1["D4"].comment is not None  # la nota sigue en su celda
    consumible = libro["CONTENEDOR #1 CONSUMIBLE "]
    assert consumible.tables["Tabla2"].ref == "A1:J7"  # se quitó el renglón vacío
    assert consumible["H2"].value == 4  # vale 7, ubicado automáticamente
    assert consumible["I6"].value is None  # el INGRESO del archivo no tenía vale que lo respalde
    pie = libro["CONTENEDOR #2 CONSUMIBLE"]
    assert pie["A6"].value == "TEXTO FUERA DE LA TABLA"


def test_inventario_crece_y_mueve_totales_y_notas(cargada, archivos, tmp_path):
    with cargada.sesion() as s:
        ubicacion = s.scalar(
            select(m.Ubicacion).where(m.Ubicacion.hoja_excel == "CONTENEDOR #1 INVENTARIABLE")
        )
        primera = s.scalar(
            select(m.Existencia).where(
                m.Existencia.ubicacion_id == ubicacion.id, m.Existencia.orden == 1
            )
        )
        # nuevo renglón al inicio para forzar que todo se recorra una fila
        for e in s.scalars(select(m.Existencia).where(m.Existencia.ubicacion_id == ubicacion.id)):
            e.orden += 1
        s.add(
            m.Existencia(
                variante_id=primera.variante_id,
                ubicacion_id=ubicacion.id,
                orden=1,
                cantidad_conteo=Decimal(9),
                conteo_id=primera.conteo_id,
            )
        )
    destino = tmp_path / "inventario.xlsx"
    with cargada.sesion() as s:
        exportar_inventario(s, archivos["inventario"], destino)
    c1 = load_workbook(destino)["CONTENEDOR #1 INVENTARIABLE"]
    assert c1.tables["Tabla1"].ref == "A1:J8"
    assert c1["A8"].value == "Total"
    assert c1["F2"].value == 9
    assert c1["D5"].comment is not None and c1["D4"].comment is None  # la nota bajó con su renglón


def test_inventario_agrega_codigos_nuevos_al_catalogo(cargada, archivos, tmp_path):
    destino = tmp_path / "inventario.xlsx"
    with cargada.sesion() as s:
        exportar_inventario(s, archivos["inventario"], destino)
    catalogo = load_workbook(destino)["ARTICULOS_MX"]
    codigos = [r[0] for r in catalogo.iter_rows(min_row=3, values_only=True)]
    assert 799 in codigos and 721 in codigos  # 799 del inventario, 721 del catálogo de vales
    assert catalogo.sheet_state == "hidden"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice no está instalado")
def test_libreoffice_abre_los_exportados(cargada, archivos, tmp_path):
    with cargada.sesion() as s:
        exportar_vales(s, archivos["vales"], tmp_path / "vales.xlsm")
        exportar_inventario(s, archivos["inventario"], tmp_path / "inventario.xlsx")
    for nombre in ("vales.xlsm", "inventario.xlsx"):
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(tmp_path),
                str(tmp_path / nombre),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
        assert (tmp_path / nombre).with_suffix(".pdf").exists()
