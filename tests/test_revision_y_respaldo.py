import datetime as dt
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from control_almacen.db import modelos as m
from control_almacen.db.sesion import BaseDatos
from control_almacen.importadores.inventario_fisico import leer_inventario
from control_almacen.importadores.vales import leer_vales
from control_almacen.respaldo.respaldo import (
    ErrorRespaldo,
    aplicar_retencion,
    crear_respaldo,
    hay_respaldo_de_hoy,
    listar_respaldos,
    restaurar_respaldo,
)
from control_almacen.servicios import revision as rv
from control_almacen.servicios.primera_carga import OpcionesPrimeraCarga, ejecutar_primera_carga
from tests.fixtures.generar import CATALOGO, FOLIO_CORTE


@pytest.fixture
def lista(archivos, tmp_path):
    vales = leer_vales(archivos["vales"])
    inventario = leer_inventario(archivos["inventario"])
    ruta = rv.generar_revision(vales, inventario, CATALOGO, tmp_path / "revision.xlsx")
    return ruta, vales


def _filas(ruta, hoja):
    return list(load_workbook(ruta)[hoja].iter_rows(min_row=2, values_only=True))


def test_revision_contiene_hallazgos(lista):
    ruta, _ = lista
    libro = load_workbook(ruta)
    assert libro["Instrucciones"]["H1"].value == rv.MARCA_FORMATO
    renglones = {f[1]: f for f in _filas(ruta, rv.HOJA_RENGLONES)}
    assert renglones[8][0] == "Alta" and renglones[8][-2] == rv.SI  # renglón perdido: eliminar
    assert renglones[6][0] == "Alta"  # encabezado perdido
    assert renglones[6][4 + 1] == "RIG 91"  # origen sugerido desde el otro renglón del folio 3
    assert "Cantidad inválida" in renglones[9][3]
    duplicados = _filas(ruta, rv.HOJA_DUPLICADOS)
    assert [(d[0], d[2], d[3], d[9]) for d in duplicados] == [(2, 4, 5, rv.SI)]
    folios = _filas(ruta, rv.HOJA_FOLIOS)
    assert folios[0][:2] == ("Alta", "8")
    codigos = _filas(ruta, rv.HOJA_CODIGOS)
    assert codigos[0][:2] == (799, "ARTICULO NUEVO")
    nombres = {f[1]: f[3] for f in _filas(ruta, rv.HOJA_NOMBRES)}
    assert nombres["MECANICO UNOO"] == "MECANICO UNO"
    valores = {(f[0], f[1]): f[3] for f in _filas(ruta, rv.HOJA_VALORES)}
    assert valores[("Destino", "(vacío)")] == "RIG 91"
    inventario = _filas(ruta, rv.HOJA_INVENTARIO)
    assert any("Renglón repetido" in f[7] for f in inventario)
    assert any(f[2] == 799 for f in inventario)


def test_lectura_de_respuestas(lista):
    ruta, vales = lista
    libro = load_workbook(ruta)
    ws = libro[rv.HOJA_RENGLONES]
    for fila in ws.iter_rows(min_row=2):
        if fila[1].value == 6:
            fila[4 + 11].value = "NUEVO"  # lote
        if fila[1].value == 9:
            fila[4 + 6].value = 15  # cantidad corregida
    ws_codigos = libro[rv.HOJA_CODIGOS]
    ws_codigos["F2"] = "ARTICULO CONFIRMADO"
    libro.save(ruta)
    respuestas = rv.leer_revision(ruta, vales.renglones)
    assert {8, 5} <= respuestas.eliminar
    assert respuestas.correcciones[6]["lote"] == "NUEVO"
    assert "cantidad" in respuestas.correcciones[9]
    assert respuestas.codigos[799] == (None, "ARTICULO CONFIRMADO")
    assert respuestas.alias["MECANICO UNOO"] == "MECANICO UNO"
    assert respuestas.normalizaciones[("destino", "(vacío)")] == "RIG 91"


def test_fila_que_ya_no_corresponde_se_ignora(lista):
    ruta, vales = lista
    libro = load_workbook(ruta)
    ws = libro[rv.HOJA_RENGLONES]
    for fila in ws.iter_rows(min_row=2):
        if fila[1].value == 6:
            fila[2].value = 999  # otro folio
    libro.save(ruta)
    respuestas = rv.leer_revision(ruta, vales.renglones)
    assert 6 not in respuestas.correcciones
    assert any("fila 6" in a for a in respuestas.advertencias)


def test_archivo_que_no_es_revision(archivos):
    with pytest.raises(rv.ErrorRevision):
        rv.leer_revision(archivos["inventario"], [])


def test_primera_carga_con_revision(bd, lista, archivos):
    ruta, vales = lista
    respuestas = rv.leer_revision(ruta, vales.renglones)
    reporte = ejecutar_primera_carga(
        bd,
        OpcionesPrimeraCarga(
            ruta_inventario=archivos["inventario"],
            ruta_vales=archivos["vales"],
            folio_corte=FOLIO_CORTE,
            fecha_conteo=dt.date(2026, 9, 3),
            respuestas=respuestas,
        ),
    )
    assert (5, "Eliminado según la revisión (duplicado o renglón inválido)") in reporte.omitidos
    with bd.sesion() as s:
        folio5 = s.scalar(select(m.Vale).where(m.Vale.folio == 5))
        assert folio5.destino == "RIG 91"
        alias = s.get(m.PersonaAlias, "MECANICO UNOO")
        assert alias is not None and alias.persona.nombre == "MECANICO UNO"


# ------------------------------------------------------------------ respaldos


@pytest.fixture
def base_con_datos(tmp_path, archivos):
    bd = BaseDatos(tmp_path / "datos" / "almacen.db")
    ejecutar_primera_carga(
        bd,
        OpcionesPrimeraCarga(
            ruta_inventario=archivos["inventario"],
            ruta_vales=archivos["vales"],
            folio_corte=FOLIO_CORTE,
            fecha_conteo=dt.date(2026, 9, 3),
            carpeta_plantillas=tmp_path / "datos" / "plantillas",
        ),
    )
    yield bd
    bd.cerrar()


def test_respaldo_y_restauracion(base_con_datos, tmp_path):
    datos = tmp_path / "datos"
    respaldos = tmp_path / "OneDrive" / "respaldos"
    ruta = crear_respaldo(base_con_datos.ruta, datos / "plantillas", respaldos, "prueba")
    with zipfile.ZipFile(ruta) as z:
        assert {"almacen.db", "manifiesto.json"} <= set(z.namelist())
        assert sum(n.startswith("plantillas/") for n in z.namelist()) == 2
    assert hay_respaldo_de_hoy(respaldos)
    # se pierde un vale y luego se restaura
    with base_con_datos.sesion() as s:
        s.delete(s.scalar(select(m.Vale).where(m.Vale.folio == 1)))
    base_con_datos.cerrar()
    nueva = tmp_path / "reinstalado" / "almacen.db"
    restaurar_respaldo(ruta, nueva, tmp_path / "reinstalado" / "plantillas")
    restaurada = BaseDatos(nueva)
    with restaurada.sesion() as s:
        assert s.scalar(select(m.Vale).where(m.Vale.folio == 1)) is not None
    restaurada.cerrar()
    assert len(list((tmp_path / "reinstalado" / "plantillas").iterdir())) == 2


def test_restaurar_archivo_invalido(tmp_path):
    malo = tmp_path / "almacen_2026-01-01_000000_x.zip"
    with zipfile.ZipFile(malo, "w") as z:
        z.writestr("otra_cosa.txt", "hola")
    with pytest.raises(ErrorRespaldo):
        restaurar_respaldo(malo, tmp_path / "x.db", tmp_path / "p")


def test_retencion(tmp_path):
    carpeta = tmp_path / "r"
    carpeta.mkdir()
    base = dt.datetime(2026, 1, 1, 8, 0, 0)
    for i in range(60):
        for j in range(2):  # dos respaldos por día
            momento = base + dt.timedelta(days=i, hours=j)
            (carpeta / f"almacen_{momento:%Y-%m-%d_%H%M%S}_auto.zip").write_bytes(b"x")
    aplicar_retencion(carpeta, diarios=30, mensuales=12)
    restantes = listar_respaldos(carpeta)
    dias = {r.fecha_hora.date() for r in restantes}
    assert len(restantes) == len(dias)  # uno por día
    # 60 días (1-ene a 1-mar): quedan los 30 más recientes (31-ene a 1-mar);
    # el último de enero ya está entre ellos, así que la retención mensual no agrega otro.
    assert len(dias) == 30
    assert min(dias) == dt.date(2026, 1, 31)
    assert restantes[0].ruta.exists()
